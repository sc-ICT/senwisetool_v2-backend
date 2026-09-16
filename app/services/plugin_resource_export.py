from __future__ import annotations

import io
import json
import zipfile
from pathlib import PurePosixPath
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.models.enums import FileNodeType
from app.models.file_node import FileNode
from app.models.plugin import PluginResource
from app.models.plugin_resource_data import PluginResourceRecord
from app.services.plugin_resource import PluginResourceService


class PluginResourceExportService:

    def __init__(
        self,
        resource_service: PluginResourceService,
    ) -> None:
        self.resource_service = resource_service

    # ======================================================================
    # EXPORT D'UNE SEULE RESSOURCE
    # ======================================================================

    async def export_resource(
        self,
        *,
        resource_id: int,
        user_id: int,
        is_admin: bool,
        storage,
    ) -> bytes:

        if not is_admin:
            raise ValueError("Seul un administrateur peut exporter " "les données d'une ressource.")

        resource = await self._load_resource(
            resource_id=resource_id,
        )

        if resource is None:
            raise ValueError("Ressource introuvable.")

        records = await self._load_records(
            resource_id=resource.id,
        )

        workbook = Workbook()

        worksheet = workbook.active

        if worksheet is None:
            worksheet = workbook.create_sheet(
                title="DATA",
            )
        else:
            worksheet.title = "DATA"

        files_to_copy: dict[str, FileNode] = {}

        await self._write_data_sheet(
            worksheet=worksheet,
            resource=resource,
            records=records,
            files_to_copy=files_to_copy,
        )

        excel_buffer = io.BytesIO()

        workbook.save(
            excel_buffer,
        )

        workbook.close()

        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(
            zip_buffer,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as archive:

            archive.writestr(
                "data.xlsx",
                excel_buffer.getvalue(),
            )

            for relative_path, file_node in files_to_copy.items():
                await self._add_file_to_archive(
                    archive=archive,
                    relative_path=relative_path,
                    file_node=file_node,
                    storage=storage,
                )

        return zip_buffer.getvalue()

    # ======================================================================
    # RESOURCE
    # ======================================================================

    async def _load_resource(
        self,
        *,
        resource_id: int,
    ) -> PluginResource | None:

        result = await self.resource_service.session.execute(
            select(PluginResource).where(
                PluginResource.id == resource_id,
            )
        )

        return result.scalar_one_or_none()

    # ======================================================================
    # DATA
    # ======================================================================

    async def _load_records(
        self,
        *,
        resource_id: int,
    ) -> list[PluginResourceRecord]:

        result = await self.resource_service.session.execute(
            select(PluginResourceRecord)
            .where(
                PluginResourceRecord.resource_id == resource_id,
            )
            .order_by(
                PluginResourceRecord.created_at.asc(),
            )
        )

        return list(
            result.scalars().all(),
        )

    # ======================================================================
    # EXCEL
    # ======================================================================

    async def _write_data_sheet(
        self,
        *,
        worksheet,
        resource: PluginResource,
        records: list[PluginResourceRecord],
        files_to_copy: dict[str, FileNode],
    ) -> None:

        fields = sorted(
            [field for field in resource.fields if field.is_active],
            key=lambda field: field.position,
        )

        # --------------------------------------------------------------
        # En-têtes
        # --------------------------------------------------------------

        for column_index, field in enumerate(
            fields,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column_index,
                value=field.key,
            )

            cell.font = Font(
                bold=True,
            )

        # --------------------------------------------------------------
        # Données
        # --------------------------------------------------------------

        for row_index, record in enumerate(
            records,
            start=2,
        ):
            data = record.data or {}

            for column_index, field in enumerate(
                fields,
                start=1,
            ):
                value = data.get(
                    field.key,
                )

                excel_value, hyperlink = await self._prepare_excel_value(
                    resource_key=resource.key,
                    record_id=record.id,
                    field_key=field.key,
                    value=value,
                    files_to_copy=files_to_copy,
                )

                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=excel_value,
                )

                if hyperlink:
                    cell.hyperlink = hyperlink

                    cell.style = "Hyperlink"

        # --------------------------------------------------------------
        # Largeur des colonnes
        # --------------------------------------------------------------

        for column_index in range(
            1,
            len(fields) + 1,
        ):
            worksheet.column_dimensions[get_column_letter(column_index)].width = 24

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions if worksheet.max_row >= 1 and worksheet.max_column >= 1 else None
        )

    # ======================================================================
    # VALUE PREPARATION
    # ======================================================================

    async def _prepare_excel_value(
        self,
        *,
        resource_key: str,
        record_id: int,
        field_key: str,
        value: Any,
        files_to_copy: dict[str, FileNode],
    ) -> tuple[Any, str | None]:

        if value is None:
            return None, None

        file_nodes = await self._find_file_nodes(
            value,
        )

        if not file_nodes:

            if isinstance(
                value,
                (dict, list),
            ):
                return (
                    json.dumps(
                        value,
                        ensure_ascii=False,
                    ),
                    None,
                )

            return value, None

        relative_paths: list[str] = []

        for file_node in file_nodes:
            safe_name = self._safe_filename(
                file_node.name,
            )

            relative_path = str(
                PurePosixPath(
                    "files",
                    resource_key,
                    str(record_id),
                    field_key,
                    safe_name,
                )
            )

            files_to_copy[relative_path] = file_node

            relative_paths.append(
                relative_path,
            )

        if not relative_paths:
            return None, None

        # --------------------------------------------------------------
        # Excel accepte un lien relatif.
        #
        # data.xlsx se trouve à la racine du ZIP.
        # Les fichiers se trouvent dans files/...
        #
        # Le lien fonctionne donc lorsque data.xlsx et files/ sont
        # extraits ensemble dans le même dossier.
        # --------------------------------------------------------------

        if len(relative_paths) == 1:
            return (
                relative_paths[0],
                relative_paths[0],
            )

        # Pour plusieurs fichiers dans la même cellule :
        # le premier est le lien cliquable et tous les chemins sont
        # conservés dans la valeur textuelle.
        return (
            " | ".join(relative_paths),
            relative_paths[0],
        )

    # ======================================================================
    # FILE NODES
    # ======================================================================

    async def _find_file_nodes(
        self,
        value: Any,
    ) -> list[FileNode]:

        ids = self._extract_file_ids(
            value,
        )

        if not ids:
            return []

        result = await self.resource_service.session.execute(
            select(FileNode).where(
                FileNode.id.in_(ids),
                FileNode.type == FileNodeType.FILE,
            )
        )

        nodes = list(
            result.scalars().all(),
        )

        by_id = {node.id: node for node in nodes}

        return [by_id[file_id] for file_id in ids if file_id in by_id]

    @classmethod
    def _extract_file_ids(
        cls,
        value: Any,
    ) -> list[int]:

        result: list[int] = []

        def visit(item: Any) -> None:

            if item is None:
                return

            if isinstance(
                item,
                int,
            ):
                result.append(
                    item,
                )
                return

            if isinstance(
                item,
                dict,
            ):

                for key, child in item.items():

                    normalized = str(key).lower()

                    if normalized in {
                        "file_id",
                        "file_node_id",
                    }:

                        if isinstance(
                            child,
                            int,
                        ):
                            result.append(
                                child,
                            )

                        continue

                    if normalized in {
                        "file_ids",
                        "file_node_ids",
                    }:

                        if isinstance(
                            child,
                            list,
                        ):

                            for child_id in child:

                                if isinstance(
                                    child_id,
                                    int,
                                ):
                                    result.append(
                                        child_id,
                                    )

                        continue

                    visit(
                        child,
                    )

                return

            if isinstance(
                item,
                list,
            ):

                for child in item:
                    visit(
                        child,
                    )

        visit(
            value,
        )

        unique: list[int] = []

        for file_id in result:

            if file_id not in unique:
                unique.append(
                    file_id,
                )

        return unique

    # ======================================================================
    # ARCHIVE
    # ======================================================================

    async def _add_file_to_archive(
        self,
        *,
        archive: zipfile.ZipFile,
        relative_path: str,
        file_node: FileNode,
        storage,
    ) -> None:

        if not file_node.storage_key:
            raise ValueError(
                f"Le fichier « {file_node.name} » " "ne possède aucune référence de stockage."
            )

        exists = await storage.exists(
            storage_key=file_node.storage_key,
        )

        if not exists:
            raise ValueError(
                f"Le fichier « {file_node.name} » "
                "est référencé par une donnée mais "
                "son contenu physique est introuvable."
            )

        source = await storage.open(
            storage_key=file_node.storage_key,
        )

        try:
            content = source.read()

            archive.writestr(
                relative_path,
                content,
            )

        finally:
            source.close()

    # ======================================================================
    # SAFE NAME
    # ======================================================================

    @staticmethod
    def _safe_filename(
        value: str,
    ) -> str:

        invalid = {
            "/",
            "\\",
            ":",
            "*",
            "?",
            '"',
            "<",
            ">",
            "|",
        }

        result = "".join("_" if char in invalid else char for char in value)

        return result or "file"
